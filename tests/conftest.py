import io
from collections.abc import Generator

import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, text
from sqlalchemy.orm import Session, sessionmaker

from app.core.config import settings
from app.core.database import Base, get_db
from app.core.security import hash_password
from app.main import app
from app.models.user import User
from app.core.enums import UserRole


TEST_DATABASE_URL = settings.DATABASE_URL.rsplit("/", 1)[0] + "/securerecruit_test"


def _ensure_test_database() -> None:
    root_url = settings.DATABASE_URL.rsplit("/", 1)[0]
    # Connect to maintenance DB. Prefer `postgres`, fallback to current DB name.
    admin_urls = [f"{root_url}/postgres", settings.DATABASE_URL]
    last_error: Exception | None = None
    for admin_url in admin_urls:
        try:
            admin_engine = create_engine(admin_url, isolation_level="AUTOCOMMIT")
            with admin_engine.connect() as conn:
                exists = conn.execute(
                    text("SELECT 1 FROM pg_database WHERE datname = 'securerecruit_test'")
                ).scalar()
                if not exists:
                    conn.execute(text("CREATE DATABASE securerecruit_test"))
            admin_engine.dispose()
            return
        except Exception as exc:  # noqa: BLE001
            last_error = exc
    raise RuntimeError(f"Could not create test database: {last_error}")


@pytest.fixture(scope="session")
def engine():
    _ensure_test_database()
    eng = create_engine(TEST_DATABASE_URL)
    Base.metadata.drop_all(bind=eng)
    Base.metadata.create_all(bind=eng)
    yield eng
    Base.metadata.drop_all(bind=eng)
    eng.dispose()


@pytest.fixture
def db(engine) -> Generator[Session, None, None]:
    connection = engine.connect()
    transaction = connection.begin()
    TestingSession = sessionmaker(bind=connection, autocommit=False, autoflush=False)
    session = TestingSession()
    try:
        yield session
    finally:
        session.close()
        transaction.rollback()
        connection.close()


@pytest.fixture
def client(db: Session) -> Generator[TestClient, None, None]:
    def override_get_db():
        try:
            yield db
        finally:
            pass

    app.dependency_overrides[get_db] = override_get_db
    with TestClient(app) as test_client:
        yield test_client
    app.dependency_overrides.clear()


@pytest.fixture
def sample_docx_bytes() -> bytes:
    buffer = io.BytesIO()
    document = Document()
    document.add_paragraph("Currículo Demo SecureRecruit")
    document.add_paragraph("Skills: Python, FastAPI, PostgreSQL, Docker, JWT")
    document.add_paragraph("Experiência: APIs REST e integração com bancos.")
    document.save(buffer)
    return buffer.getvalue()


def create_user(db: Session, *, email: str, role: str, name: str = "User Demo") -> User:
    user = User(
        name=name,
        email=email,
        password_hash=hash_password("Demo@1234"),
        role=role,
        is_active=True,
    )
    db.add(user)
    db.commit()
    db.refresh(user)
    return user


def auth_header(client: TestClient, email: str, password: str = "Demo@1234") -> dict[str, str]:
    response = client.post("/api/v1/auth/login", json={"email": email, "password": password})
    assert response.status_code == 200, response.text
    token = response.json()["access_token"]
    return {"Authorization": f"Bearer {token}"}
